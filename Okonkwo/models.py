from django.db import models
from django.core.exceptions import ValidationError

import os
import pymupdf

from docx import Document
from pptx import Presentation



# OLD PUBLICATION VALIDATOR


def validate_publication_file(value):
    """
    Kept for compatibility with existing Django migrations.

    Publication files are no longer restricted by this validator.
    """
    return



# BLOG POST FILE VALIDATION


def validate_blog_file(value):
    """
    Allows:

    - Normal image files
    - PDF files containing at least one image
    - DOCX files containing at least one image
    - PPTX files containing at least one image

    Documents without images are rejected.
    """

    file_name = value.name.lower()
    extension = os.path.splitext(file_name)[1]

    allowed_image_extensions = {
        ".jpg",
        ".jpeg",
        ".png",
        ".gif",
        ".webp",
        ".bmp",
        ".tiff",
        ".tif",
    }

    # =====================================================
    # NORMAL IMAGE
    # =====================================================

    if extension in allowed_image_extensions:
        return

    # =====================================================
    # PDF
    # =====================================================

    if extension == ".pdf":

        try:
            value.seek(0)

            pdf_data = value.read()

            value.seek(0)

            document = pymupdf.open(
                stream=pdf_data,
                filetype="pdf"
            )

            contains_image = False

            for page in document:

                images = page.get_images(
                    full=True
                )

                if images:
                    contains_image = True
                    break

            document.close()

            if not contains_image:

                raise ValidationError(
                    "This PDF cannot be uploaded because it does "
                    "not contain an image."
                )

            return

        except ValidationError:
            raise

        except Exception:

            raise ValidationError(
                "The PDF could not be checked. "
                "Please upload a valid PDF containing an image."
            )

    # =====================================================
    # DOCX
    # =====================================================

    if extension == ".docx":

        try:

            value.seek(0)

            document = Document(
                value
            )

            value.seek(0)

            image_count = 0

            for relationship in document.part.rels.values():

                if "image" in relationship.reltype:

                    image_count += 1

            if image_count == 0:

                raise ValidationError(
                    "This Word document cannot be uploaded because "
                    "it does not contain an image."
                )

            return

        except ValidationError:
            raise

        except Exception:

            raise ValidationError(
                "The Word document could not be checked. "
                "Please upload a valid DOCX document containing an image."
            )

    # =====================================================
    # PPTX
    # =====================================================

    if extension == ".pptx":

        try:

            value.seek(0)

            presentation = Presentation(
                value
            )

            value.seek(0)

            contains_image = False

            for slide in presentation.slides:

                for shape in slide.shapes:

                    if shape.shape_type == 13:

                        contains_image = True

                        break

                if contains_image:
                    break

            if not contains_image:

                raise ValidationError(
                    "This PowerPoint presentation cannot be uploaded "
                    "because it does not contain an image."
                )

            return

        except ValidationError:
            raise

        except Exception:

            raise ValidationError(
                "The PowerPoint presentation could not be checked. "
                "Please upload a valid PPTX presentation containing an image."
            )

    # =====================================================
    # EVERYTHING ELSE
    # =====================================================

    raise ValidationError(
        "Only image files, PDF documents, Word documents, "
        "and PowerPoint presentations containing images are allowed."
    )



# TEAM MEMBERS


class TeamMember(models.Model):

    # These are suggested categories.
    # The admin will allow users to type a different category too.

    CATEGORY_CHOICES = [
        ("director", "Director"),
        ("professor", "Professor"),
        ("postdoc", "Postdoctoral Researcher"),
        ("phd", "PhD Student"),
        ("masters", "Master's Student"),
        ("undergraduate", "Undergraduate Student"),
        ("research_assistant", "Research Assistant"),
        ("visiting_scholar", "Visiting Scholar"),
        ("visiting_student", "Visiting Student"),
        ("alumni", "Alumni"),
        ("other", "Other"),
    ]

    name = models.CharField(
        max_length=200
    )

    position = models.CharField(
        max_length=200
    )

    # NOT a fixed Django choices field.
    # Users can select a suggested category or type their own.

    category = models.CharField(
        max_length=100,
        default="other",
        verbose_name="Category"
    )

    photo = models.ImageField(
        upload_to="team/",
        blank=True,
        null=True
    )

    email = models.EmailField(
        blank=True
    )

    linkedin = models.URLField(
        blank=True
    )

    order = models.PositiveIntegerField(
        default=0
    )

    class Meta:
        ordering = [
            "order",
            "name"
        ]

    def __str__(self):
        return self.name



# PUBLICATIONS


class Publication(models.Model):

    PUBLICATION_TYPES = [
        ("journal", "Journal Article"),
        ("conference", "Conference Paper"),
        ("book", "Book / Book Chapter"),
        ("report", "Research Report"),
        ("working", "Working Paper"),
        ("other", "Other"),
    ]

    title = models.CharField(
        max_length=500
    )

    authors = models.CharField(
        max_length=500
    )

    journal = models.CharField(
        max_length=300,
        blank=True
    )

    year = models.PositiveIntegerField()

    publication_type = models.CharField(
        max_length=20,
        choices=PUBLICATION_TYPES,
        default="journal"
    )

    abstract = models.TextField(
        blank=True
    )

    image = models.ImageField(
        upload_to="publications/",
        blank=True,
        null=True
    )

    pdf_file = models.FileField(
        upload_to="publications/pdfs/",
        blank=True,
        null=True,
        verbose_name="Full Text File"
    )

    download_count = models.PositiveIntegerField(
        default=0,
        editable=False
    )

    journal_url = models.URLField(
        blank=True,
        verbose_name="Journal Link"
    )

    google_scholar_url = models.URLField(
        blank=True,
        verbose_name="Google Scholar Link"
    )

    researchgate_url = models.URLField(
        blank=True,
        verbose_name="ResearchGate Link"
    )

    doi = models.URLField(
        blank=True,
        verbose_name="DOI Link"
    )

    publication_url = models.URLField(
        blank=True,
        verbose_name="Publication Link"
    )

    featured = models.BooleanField(
        default=False
    )

    date_uploaded = models.DateTimeField(
        auto_now_add=True,
        verbose_name="Date of Upload"
    )

    class Meta:
        ordering = [
            "-year",
            "title"
        ]

    def __str__(self):
        return self.title



# UPDATE PROFILE


class UpdateProfile(models.Model):

    name = models.CharField(
        max_length=200
    )

    title = models.CharField(
        max_length=300
    )

    department = models.CharField(
        max_length=300,
        blank=True
    )

    institution = models.CharField(
        max_length=300,
        blank=True
    )

    email = models.EmailField(
        blank=True
    )

    photo = models.ImageField(
        upload_to="profile/",
        blank=True,
        null=True
    )

    bio = models.TextField(
        blank=True
    )

    google_scholar = models.URLField(
        blank=True
    )

    live_DNA = models.URLField(
        blank=True
    )

    linkedin = models.URLField(
        blank=True
    )

    ORCID = models.URLField(
        blank=True
    )

    ResearchGate = models.URLField(
        blank=True
    )

    def __str__(self):
        return self.name



# BLOG POSTS


class BlogPost(models.Model):

    post_title = models.CharField(
        max_length=300,
        verbose_name="Post Title"
    )

    label = models.CharField(
        max_length=100
    )

    abstract = models.TextField(
        blank=True,
        verbose_name="Abstract"
    )

    # =====================================================
    # ONE FILE FIELD
    #
    # Allows:
    # - Images
    # - PDF containing an image
    # - DOCX containing an image
    # - PPTX containing an image
    # =====================================================

    header_picture = models.FileField(
        upload_to="blog/",
        blank=True,
        null=True,
        verbose_name="Post File",
        validators=[
            validate_blog_file
        ]
    )

    post_count = models.PositiveIntegerField(
        default=0,
        editable=False,
        verbose_name="Post Count"
    )

    date_uploaded = models.DateTimeField(
        auto_now_add=True,
        verbose_name="Date of Upload"
    )

    class Meta:
        ordering = [
            "-date_uploaded"
        ]

        verbose_name = "Blog Post"

        verbose_name_plural = "Blog Posts"

    def __str__(self):
        return self.post_title



# BLOG COMMENTS


class BlogComment(models.Model):

    blog_post = models.ForeignKey(
        BlogPost,
        on_delete=models.CASCADE,
        related_name="comments",
        verbose_name="Blog Post"
    )

    name = models.CharField(
        max_length=150,
        verbose_name="Name"
    )

    email = models.EmailField(
        verbose_name="Email"
    )

    comment = models.TextField(
        verbose_name="Comment"
    )

    approved = models.BooleanField(
        default=False,
        verbose_name="Approved"
    )

    date_commented = models.DateTimeField(
        auto_now_add=True,
        verbose_name="Date Commented"
    )

    class Meta:
        ordering = [
            "date_commented"
        ]

        verbose_name = "Blog Comment"

        verbose_name_plural = "Blog Comments"

    def __str__(self):
        return f"{self.name} - {self.blog_post.post_title}"